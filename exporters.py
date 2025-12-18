# exporters.py
from __future__ import annotations

import csv
import json
import os
from datetime import datetime
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


def _safe_stem(s: str) -> str:
    return "".join(c for c in (s or "output") if c.isalnum() or c in ("-", "_")).strip("_")


def ensure_outdir(outdir: str) -> str:
    os.makedirs(outdir, exist_ok=True)
    return outdir


def _autosize_worksheet(ws, max_width: int = 70) -> None:
    # Simple autosize based on cell string length
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is None:
                continue
            s = str(cell.value)
            if len(s) > max_len:
                max_len = len(s)
        ws.column_dimensions[col_letter].width = min(max_width, max(10, max_len + 2))


def _write_header_row(ws, headers: List[str]) -> None:
    bold = Font(bold=True)
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = bold
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def export_quote_bundle(result: Dict[str, Any], outdir: str, stem: Optional[str] = None) -> Dict[str, str]:
    """
    Writes:
      - *_quote.json
      - *_route.csv
      - *_route_times.csv
      - *_summary.txt
      - *_route_times.xlsx  (NEW: includes minutes+hours + totals)
    Returns dict of generated file paths.
    """
    outdir = ensure_outdir(outdir)

    drawing_no = result.get("drawing_no") or "UNKNOWN"
    title = result.get("title") or ""
    if stem is None:
        stem = f"{drawing_no}_{_safe_stem(title) or 'quote'}"

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = os.path.join(outdir, f"{stem}_{ts}")

    paths: Dict[str, str] = {}

    route = result.get("route") or []
    totals = result.get("totals") or {}
    env = result.get("envelope_mm") or {}
    notes = result.get("notes") or []
    warnings = result.get("warnings") or []

    # 1) Full JSON
    json_path = base + "_quote.json"
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(result, f, indent=2)
    paths["quote_json"] = json_path

    # 2) Route CSV (ops only)
    route_csv = base + "_route.csv"
    with open(route_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["op_no", "operation", "machine_group", "description"])
        w.writeheader()
        for op in route:
            w.writerow({
                "op_no": op.get("op_no"),
                "operation": op.get("operation"),
                "machine_group": op.get("machine_group"),
                "description": op.get("description"),
            })
    paths["route_csv"] = route_csv

    # 3) Route times CSV (mins only + total_min)
    route_times_csv = base + "_route_times.csv"
    with open(route_times_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(
            f,
            fieldnames=["op_no", "operation", "machine_group", "setup_min", "runtime_min", "total_min", "description"],
        )
        w.writeheader()
        for op in route:
            w.writerow({
                "op_no": op.get("op_no"),
                "operation": op.get("operation"),
                "machine_group": op.get("machine_group"),
                "setup_min": op.get("setup_min"),
                "runtime_min": op.get("runtime_min"),
                "total_min": op.get("total_min"),
                "description": op.get("description"),
            })
    paths["route_times_csv"] = route_times_csv

    # 4) Summary TXT
    summary_path = base + "_summary.txt"
    with open(summary_path, "w", encoding="utf-8") as f:
        f.write(f"Drawing: {drawing_no}\n")
        f.write(f"Title: {title}\n")
        f.write(f"Material: {result.get('material')}\n")
        f.write(f"Weight (kg): {result.get('finished_weight_kg')}\n\n")
        f.write("Envelope (mm):\n")
        f.write(f"  L x W x T = {env.get('length_mm')} x {env.get('width_mm')} x {env.get('thickness_mm')}\n\n")
        f.write("Totals:\n")
        f.write(f"  Total minutes: {totals.get('total_min')}\n")
        f.write(f"  Total hours:   {totals.get('total_hr')}\n\n")
        f.write("Notes:\n")
        for n in notes:
            f.write(f"  - {n}\n")
        if not notes:
            f.write("  (none)\n")
        f.write("\nWarnings:\n")
        for wmsg in warnings:
            f.write(f"  - {wmsg}\n")
        if not warnings:
            f.write("  (none)\n")
    paths["summary_txt"] = summary_path

    # 5) XLSX (NEW)
    xlsx_path = base + "_route_times.xlsx"
    wb = Workbook()

    # --- Summary sheet ---
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_sum.append(["Drawing No", drawing_no])
    ws_sum.append(["Title", title])
    ws_sum.append(["Material", result.get("material")])
    ws_sum.append(["Weight (kg)", result.get("finished_weight_kg")])
    ws_sum.append(["Dims source", result.get("dims_source")])
    ws_sum.append(["", ""])
    ws_sum.append(["Envelope L (mm)", env.get("length_mm")])
    ws_sum.append(["Envelope W (mm)", env.get("width_mm")])
    ws_sum.append(["Envelope T (mm)", env.get("thickness_mm")])
    ws_sum.append(["", ""])
    ws_sum.append(["TOTAL minutes", totals.get("total_min")])
    ws_sum.append(["TOTAL hours", totals.get("total_hr")])

    ws_sum.append(["", ""])
    ws_sum.append(["Notes", ""])
    if notes:
        for n in notes:
            ws_sum.append(["", n])
    else:
        ws_sum.append(["", "(none)"])

    ws_sum.append(["", ""])
    ws_sum.append(["Warnings", ""])
    if warnings:
        for wmsg in warnings:
            ws_sum.append(["", wmsg])
    else:
        ws_sum.append(["", "(none)"])

    # Format summary
    bold = Font(bold=True)
    for r in range(1, ws_sum.max_row + 1):
        ws_sum.cell(row=r, column=1).font = bold
        ws_sum.cell(row=r, column=1).alignment = Alignment(vertical="top")
        ws_sum.cell(row=r, column=2).alignment = Alignment(vertical="top", wrap_text=True)
    _autosize_worksheet(ws_sum, max_width=90)

    # --- Route_Times sheet ---
    ws_rt = wb.create_sheet("Route_Times", 1)
    headers = [
        "Op No", "Operation", "Machine Group",
        "Setup (min)", "Runtime (min)", "Total (min)",
        "Setup (hr)", "Runtime (hr)", "Total (hr)",
        "Description"
    ]
    _write_header_row(ws_rt, headers)

    def _to_float(x) -> float:
        try:
            return float(x)
        except Exception:
            return 0.0

    total_setup_min = 0.0
    total_runtime_min = 0.0
    total_total_min = 0.0

    for op in route:
        setup_min = _to_float(op.get("setup_min"))
        runtime_min = _to_float(op.get("runtime_min"))
        total_min = _to_float(op.get("total_min"))

        total_setup_min += setup_min
        total_runtime_min += runtime_min
        total_total_min += total_min

        ws_rt.append([
            op.get("op_no"),
            op.get("operation"),
            op.get("machine_group"),
            setup_min,
            runtime_min,
            total_min,
            setup_min / 60.0,
            runtime_min / 60.0,
            total_min / 60.0,
            op.get("description"),
        ])

    # Totals row
    ws_rt.append(["", "", "TOTALS",
                  total_setup_min, total_runtime_min, total_total_min,
                  total_setup_min / 60.0, total_runtime_min / 60.0, total_total_min / 60.0,
                  ""])

    # Format numeric columns
    for row in ws_rt.iter_rows(min_row=2, max_row=ws_rt.max_row):
        # minutes cols 4-6
        for c in row[3:6]:
            c.number_format = "0.00"
        # hours cols 7-9
        for c in row[6:9]:
            c.number_format = "0.000"
        # wrap description
        row[9].alignment = Alignment(wrap_text=True, vertical="top")

    # Make totals row bold
    for c in ws_rt[ws_rt.max_row]:
        c.font = Font(bold=True)

    _autosize_worksheet(ws_rt, max_width=90)

    wb.save(xlsx_path)
    paths["route_times_xlsx"] = xlsx_path

    return paths
